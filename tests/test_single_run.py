"""The functions both execution shapes share, tested without a browser.

The three-step shape and the single task call the same code to start a run,
read a row and finish a run. These tests pin that code down, and the two rules
that only the single task relies on: a list with no usable link never starts a
browser, and a list with many links signs in once. Every value is invented.
"""

from datetime import date, datetime, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import tasks
from vigiflow_tasks.email_run import RunContext
from vigiflow_tasks.mailer import MailError
from vigiflow_tasks.report_scraper import ScrapedReport

GUID_A = "e71a7f06-ba16-415e-a827-134939a6323e"
GUID_B = "6142776a-503d-45c3-8649-89f0c0d0a6c0"
URL = "https://vigiflow.who-umc.org/dataentry/{}"


class FakeItem:
    """A trigger work item with no parsed email, as a run started by hand has."""

    def __init__(self, payload):
        self.payload = payload
        self.files = []
        self.failures = []

    def email(self, html=False):
        raise ValueError("no email on this item")

    def fail(self, exception_type, code=None, message=None):
        self.failures.append((exception_type, code))


class FakeClient:
    """A signed-in session that counts how often it was asked to sign in."""

    logins = 0

    def __init__(self):
        self.page = object()

    def __enter__(self):
        return self

    def __exit__(self, *_):
        return False

    def login(self):
        FakeClient.logins += 1


def _links(tmp_path, rows):
    book = Workbook()
    book.active.append(["Links"])
    for row in rows:
        book.active.append([row])
    path = tmp_path / "Links.xlsx"
    book.save(path)
    return path


def _report(url):
    return ScrapedReport(
        url=url,
        report_id="PE-DIGEMID-399999001",
        fields={
            "fecha_recepcion_reciente": date(2026, 9, 3),
            "fecha_notificacion": None,
            "eess": "CENTRO DE SALUD EJEMPLO",
            "tipo_reporte": "RAM",
            "ev": None,
            "id_codigo": "PE-DIGEMID-399999001",
            "codigo_ipress": "P-01 ABC EJEMPLO RAM",
            "paciente": "A.B.C",
            "medicamentos_sospechosos": ["MEDICAMENTO A"],
            "gravedad": "LEVE",
            "_clinic_assessed": False,
        },
    )


@pytest.fixture(autouse=True)
def _quiet_env(monkeypatch):
    monkeypatch.delenv("VIGIFLOW_LINKS_FILE", raising=False)
    monkeypatch.delenv("VIGIFLOW_SMTP_REPLY_TO", raising=False)
    monkeypatch.delenv("VIGIFLOW_TIMEZONE", raising=False)
    monkeypatch.delenv("VIGIFLOW_VALIDATION_DATE", raising=False)
    FakeClient.logins = 0


# -- starting a run ----------------------------------------------------------


def test_a_run_starts_with_every_row_in_spreadsheet_order(tmp_path):
    path = _links(tmp_path, [URL.format(GUID_A), "N/A", URL.format(GUID_B)])
    item = FakeItem({"reply_to": "analista@example.org", "links_file": str(path)})

    context, queued = tasks._start_run(item, tmp_path / "out", datetime.now(timezone.utc))

    assert item.failures == []
    assert context.reply_to == "analista@example.org"
    assert context.total_links == 3
    assert [q["source_row"] for q in queued] == [2, 3, 4]
    assert [q["position"] for q in queued] == [1, 2, 3]
    assert queued[0]["url"].endswith(GUID_A)
    assert queued[1]["problem"] and "url" not in queued[1]
    assert all(q["run_id"] == context.run_id for q in queued)


def test_a_trigger_without_a_spreadsheet_is_a_business_failure(tmp_path):
    item = FakeItem({"reply_to": "analista@example.org"})
    assert tasks._start_run(item, tmp_path, datetime.now(timezone.utc)) is None
    assert item.failures == [("BUSINESS", "NO_ATTACHMENT")]


def test_a_trigger_without_a_sender_is_a_business_failure(tmp_path):
    item = FakeItem({"links_file": str(_links(tmp_path, [URL.format(GUID_A)]))})
    assert tasks._start_run(item, tmp_path, datetime.now(timezone.utc)) is None
    assert item.failures == [("BUSINESS", "NO_REPLY_ADDRESS")]


# -- reading a row -----------------------------------------------------------


def test_a_cell_that_is_not_a_link_never_touches_the_browser(tmp_path, monkeypatch):
    monkeypatch.setattr(tasks, "scrape_report", lambda *_: pytest.fail("opened a page"))
    tally = tasks.Tally()
    outcome = tasks._read_row(None, {"source_row": 3, "problem": "no address", "value": "N/A"}, tmp_path, tally)

    assert outcome["broken"] and outcome["code"] == tasks.NOT_A_LINK
    assert outcome["screenshot_path"] is None
    assert tally.broken == 1


def test_a_link_that_is_not_a_report_is_photographed(tmp_path, monkeypatch):
    def fake_screenshot(_page, destination):
        Path(destination).write_bytes(b"png")
        return Path(destination)

    monkeypatch.setattr(tasks, "scrape_report", lambda _page, url: ScrapedReport(url=url))
    monkeypatch.setattr(tasks, "screenshot", fake_screenshot)
    tally = tasks.Tally()

    outcome = tasks._read_row(object(), {"source_row": 4, "url": URL.format(GUID_A)}, tmp_path, tally)

    assert outcome["code"] == tasks.LINK_NOT_A_REPORT
    assert outcome["screenshot_path"].endswith("enlace-fila-4.png")
    assert tasks._broken_message(outcome) == "Row 4 does not lead to a report"


def test_a_report_comes_back_ready_for_the_queue(tmp_path, monkeypatch):
    monkeypatch.setattr(tasks, "scrape_report", lambda _page, url: _report(url))
    tally = tasks.Tally()

    outcome = tasks._read_row(object(), {"source_row": 2, "position": 1, "url": URL.format(GUID_A)}, tmp_path, tally)

    assert outcome["report_id"] == "PE-DIGEMID-399999001"
    assert outcome["fields"]["fecha_recepcion_reciente"] == "2026-09-03"  # survives JSON
    assert tally.read == 1


# -- the single task's reading loop ------------------------------------------


def test_a_list_with_no_usable_link_never_starts_a_browser(tmp_path, monkeypatch):
    monkeypatch.setattr(tasks, "open_client", lambda *_: pytest.fail("started a browser"))
    queued = [{"source_row": 2, "problem": "no address"}, {"source_row": 3, "problem": "no address"}]

    outcomes = tasks._read_rows(queued, tmp_path, tasks.Tally())

    assert [o["broken"] for o in outcomes] == [True, True]


def test_many_links_are_read_in_one_session_with_one_sign_in(tmp_path, monkeypatch):
    monkeypatch.setattr(tasks, "open_client", lambda *_: FakeClient())
    monkeypatch.setattr(tasks, "scrape_report", lambda _page, url: _report(url))
    queued = [{"source_row": n, "position": n - 1, "url": URL.format(GUID_A)} for n in range(2, 12)]

    outcomes = tasks._read_rows(queued, tmp_path, tasks.Tally())

    assert len(outcomes) == 10
    assert FakeClient.logins == 1


# -- finishing a run ---------------------------------------------------------


def _context():
    return RunContext(
        run_id="abc123", reply_to="analista@example.org", input_file="Links.xlsx",
        started_at=datetime(2026, 9, 14, 13, 0, tzinfo=timezone.utc).isoformat(), total_links=2,
    )


def _finished_inputs(tmp_path, monkeypatch):
    monkeypatch.setattr(tasks, "scrape_report", lambda _page, url: _report(url))
    row = tasks._read_row(object(), {"source_row": 2, "position": 1, "url": URL.format(GUID_A)},
                          tmp_path, tasks.Tally())
    shot = tmp_path / "enlace-fila-3.png"
    shot.write_bytes(b"png")
    (tmp_path / "Links.xlsx").write_bytes(b"list")
    broken = [tasks._reply_entry({"source_row": 3, "position": 2, "value": "x", "detail": "d"}, str(shot))]
    return [row], broken, shot


def test_one_reply_carries_the_workbook_and_screenshots_then_the_files_go(tmp_path, monkeypatch):
    rows, broken, shot = _finished_inputs(tmp_path, monkeypatch)
    sent = []
    monkeypatch.setattr(
        tasks, "send_report",
        lambda to, subject, body, attachments=None, **_: sent.append((to, [Path(a) for a in attachments])) or to,
    )

    tasks._finish_run(_context(), rows, broken, tmp_path, datetime.now(timezone.utc))

    assert len(sent) == 1
    to, attachments = sent[0]
    assert to == "analista@example.org"
    workbook = next(a for a in attachments if a.suffix == ".xlsx")
    assert workbook.name.startswith("REPORTES DEL VIGIFLOW DEL DIA")
    assert shot.name in [a.name for a in attachments]
    assert not workbook.exists() and not shot.exists() and not (tmp_path / "Links.xlsx").exists()


def test_a_failed_send_deletes_nothing(tmp_path, monkeypatch):
    rows, broken, shot = _finished_inputs(tmp_path, monkeypatch)

    def refuse(*_, **__):
        raise MailError("server said no")

    monkeypatch.setattr(tasks, "send_report", refuse)
    monkeypatch.setattr(tasks, "_diagnose_mail", lambda: None)

    with pytest.raises(MailError):
        tasks._finish_run(_context(), rows, broken, tmp_path, datetime.now(timezone.utc))

    workbooks = list(tmp_path.glob("REPORTES*.xlsx"))
    assert len(workbooks) == 1 and shot.exists()
    assert load_workbook(workbooks[0]).active.max_row >= 2

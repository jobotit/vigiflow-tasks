"""Builds the empty Excel template from a filled daily report.

The business gave us a finished file rather than a blank template, so the
template is that file with the data rows removed. Header text, fill, borders,
row height, column widths and the date number formats are all preserved, so
the robot's output is visually identical to what is produced by hand today.

Usage:
    python scripts/make_template_from_sample.py "path/to/REPORTES ... .xlsx"

Re-run this if the business changes the report layout.
"""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

DEFAULT_SAMPLE = Path.home() / "Downloads" / "REPORTES DEL VIGIFLOW DEL DÍA 31-08-2026.xlsx"
DESTINATION = Path("resources/templates/reportes_vigiflow.xlsx")
FIRST_DATA_ROW = 2


def main() -> int:
    from openpyxl import load_workbook

    sample = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SAMPLE
    if not sample.exists():
        print(f"Sample not found: {sample}")
        return 1

    DESTINATION.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(sample, DESTINATION)

    workbook = load_workbook(DESTINATION)
    sheet = workbook.active

    # Keep one styled data row as the style carrier, then delete the rest.
    # openpyxl loses per-cell styling on rows that never existed, so the
    # writer copies the style of row 2 onto each row it adds.
    last_row = sheet.max_row
    if last_row > FIRST_DATA_ROW:
        sheet.delete_rows(FIRST_DATA_ROW + 1, last_row - FIRST_DATA_ROW)

    # Blank the remaining data row's values but keep its formatting.
    for column in range(1, sheet.max_column + 1):
        sheet.cell(row=FIRST_DATA_ROW, column=column).value = None

    workbook.save(DESTINATION)

    print(f"Template written: {DESTINATION}")
    print(f"  sheet      : {sheet.title}")
    print(f"  columns    : {sheet.max_column}")
    print("  header row : 1")
    print(f"  data starts: {FIRST_DATA_ROW}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

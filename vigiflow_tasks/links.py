"""Reading the analyst's list of reports from a spreadsheet.

The file is a single column of links:

    Links
    https://vigiflow.who-umc.org/dataentry/<guid>
    ...

Every non-blank row becomes a transaction, including the rows that are not
usable links. That is deliberate. A row quietly dropped here is a report the
analyst believes was processed and was not, so a row that cannot be used is
queued as a problem and reported back by email rather than skipped.

Rows fall into three kinds:

    a report link   on the VigiFlow host. Opened and read normally.

    openable        on the VigiFlow host but not a report address. Queued
                    anyway, so the consumer opens it and can show, in a
                    screenshot, what is actually there.

    not a link      anything else: a stray note, an address for some other
                    system, an empty-looking cell with a stray character.
                    There is nothing to open, so it is reported without one.

Duplicates are a fourth case and not a problem: the same report listed twice
is read once, and the repeat is noted rather than reported as an error.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from urllib.parse import urlparse

from vigiflow_tasks.vigiflow import locators as loc

LOGGER = logging.getLogger(__name__)

GUID_PATTERN = re.compile(
    r"[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}", re.IGNORECASE
)
HEADER_NAMES = ("links", "link", "url", "enlace", "enlaces")

# What the analyst is told about a row that could not even be opened. The
# wording matches the one the consumer uses for a link that opened but turned
# out not to be a report, because to the person reading the email they are the
# same problem.
NOT_A_REPORT = "The link provided is broken or is not related to a report."
NOT_A_LINK_DETAIL = "The cell is not a VigiFlow address, so there was nothing to open."
NOT_A_REPORT_DETAIL = "The page it opens is not a report."

# How much of an unusable cell is quoted back. Enough to recognise the row,
# short enough that a cell holding something long does not fill the email.
QUOTE_LIMIT = 120


class LinksError(RuntimeError):
    """The list of links could not be read at all."""


@dataclass
class LinkedReport:
    """One row of the analyst's list that can be opened."""

    url: str
    guid: str
    row: int


@dataclass
class LinkProblem:
    """One row of the analyst's list that cannot be opened."""

    row: int
    value: str
    detail: str = NOT_A_LINK_DETAIL


@dataclass
class LinkList:
    path: str
    reports: list[LinkedReport] = field(default_factory=list)
    problems: list[LinkProblem] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)

    @property
    def rows(self) -> int:
        """Every row that became a transaction, usable or not."""
        return len(self.reports) + len(self.problems)

    def describe(self) -> str:
        text = f"{len(self.reports)} links from {Path(self.path).name}"
        if self.problems:
            text += f", {len(self.problems)} rows that are not links"
        if self.skipped:
            text += f", {len(self.skipped)} rows skipped"
        return text


def read_links(path: Path | str) -> LinkList:
    """Read the analyst's list.

    Accepts any column whose header looks like a link column, and falls back
    to the first column when none does, so a file saved without a header still
    works.

    Raises only when the file itself cannot be used. A file full of unusable
    rows is not an error here: those rows are the problems the analyst needs
    told about, and dropping them would be the silent failure this is meant to
    avoid.
    """
    from openpyxl import load_workbook

    path = Path(path)
    if not path.exists():
        raise LinksError(f"List of links not found: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as err:
        raise LinksError(f"Could not open {path.name}: {err}") from err

    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook.close()
    if not rows:
        raise LinksError(f"{path.name} is empty")

    column = _link_column(rows[0])
    start = 2 if _looks_like_header(rows[0], column) else 1

    result = LinkList(path=str(path))
    seen: set[str] = set()

    for number, row in enumerate(rows[start - 1:], start=start):
        value = row[column] if column < len(row) else None
        text = str(value or "").strip()
        if not text:
            continue

        url = _as_report_address(text)
        if url is None:
            result.problems.append(
                LinkProblem(row=number, value=text[:QUOTE_LIMIT], detail=NOT_A_LINK_DETAIL)
            )
            continue

        match = GUID_PATTERN.search(url)
        guid = match.group(0).lower() if match else ""
        # A report is keyed by its guid; an address without one is deduplicated
        # by the address itself, since two spellings of it are two chances to
        # open the same wrong page.
        key = guid or url.rstrip("/").lower()
        if key in seen:
            result.skipped.append(f"row {number}: duplicate of an earlier link")
            continue
        seen.add(key)
        result.reports.append(LinkedReport(url=url, guid=guid, row=number))

    if not result.rows:
        raise LinksError(f"{path.name} has no rows with anything in them")

    LOGGER.info("Read %s", result.describe())
    return result


def _as_report_address(text: str) -> str | None:
    """The address to open for this cell, or None if there is nothing to open.

    A cell holding a bare guid and no address is accepted and turned into one,
    because an analyst pasting from the id column produces exactly that.

    Otherwise the host is checked, but not the path. A link to the right system
    with the wrong path is exactly the case where a screenshot tells the
    analyst something, so it is opened rather than rejected here. Anything
    pointing somewhere else is not opened at all: this process has no business
    browsing to an arbitrary address that arrived by email.
    """
    text = text.strip()
    if GUID_PATTERN.fullmatch(text):
        return loc.REPORT_URL_TEMPLATE.format(host=loc.APP_HOST, guid=text.lower())
    try:
        parsed = urlparse(text)
    except ValueError:
        return None
    if parsed.scheme not in ("http", "https"):
        return None
    if (parsed.hostname or "").lower() != loc.APP_HOST.lower():
        return None
    return text


def _link_column(header_row) -> int:
    """Which column holds the links."""
    for index, cell in enumerate(header_row or ()):
        if str(cell or "").strip().lower() in HEADER_NAMES:
            return index
    # No recognisable header, so assume the first column that holds a GUID.
    for index, cell in enumerate(header_row or ()):
        if GUID_PATTERN.search(str(cell or "")):
            return index
    return 0


def _looks_like_header(row, column: int) -> bool:
    value = str(row[column] if column < len(row) else "" or "").strip()
    return bool(value) and not GUID_PATTERN.search(value)

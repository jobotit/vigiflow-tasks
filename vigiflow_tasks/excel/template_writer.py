"""Fills the daily VigiFlow report workbook from an Excel template.

The template is never edited in place: each batch gets a fresh copy, so a
failed run cannot corrupt the template and reruns are safe.

The layout is described by ``resources/field_mapping.json``, keyed by column
letter. Each column declares where its value comes from:

``field``
    Read from VigiFlow for this report.
``list``
    One VigiFlow field holding several values, spread one per column. The
    five "medicamento sospechoso" columns work this way.
``sequence``
    The running correlative (column N.). Assigned by the producer so that
    batches processed out of order still number consecutively.
``run``
    Constant for the whole run, such as the validation date. Supplied by the
    task rather than read from VigiFlow.
"""

from __future__ import annotations

import json
import logging
import re
import shutil
from copy import copy
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from vigiflow_tasks.config import FIELD_MAPPING_PATH
from vigiflow_tasks.models import ReportRecord

LOGGER = logging.getLogger(__name__)

FIELD_SOURCES = ("field", "list", "sequence", "run")
VALUE_TYPES = ("text", "date", "int", "digits", "auto")
DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y", "%d.%m.%Y")


class TemplateError(RuntimeError):
    """The template or its mapping cannot be used."""


@dataclass
class ColumnSpec:
    """One column of the report sheet."""

    letter: str
    source: str
    name: str
    type: str = "text"
    index: int = 0
    header: str | None = None

    @classmethod
    def parse(cls, letter: str, raw: dict[str, Any]) -> "ColumnSpec":
        source = raw.get("source", "field")
        if source not in FIELD_SOURCES:
            raise TemplateError(
                f"Column {letter}: source must be one of {', '.join(FIELD_SOURCES)}, got {source!r}"
            )
        value_type = raw.get("type", "text")
        if value_type not in VALUE_TYPES:
            raise TemplateError(
                f"Column {letter}: type must be one of {', '.join(VALUE_TYPES)}, got {value_type!r}"
            )
        name = raw.get("name")
        if not name:
            raise TemplateError(f"Column {letter}: 'name' is required")
        return cls(
            letter=letter,
            source=source,
            name=str(name),
            type=value_type,
            index=int(raw.get("index", 0)),
            header=raw.get("header"),
        )


@dataclass
class FieldMapping:
    """Which report field goes into which column of the sheet."""

    sheet: str | None
    header_row: int
    start_row: int
    style_row: int | None
    columns: list[ColumnSpec] = field(default_factory=list)

    @classmethod
    def load(cls, path: Path | None = None) -> "FieldMapping":
        path = Path(path or FIELD_MAPPING_PATH)
        if not path.exists():
            raise TemplateError(f"Field mapping not found: {path}")
        data = json.loads(path.read_text(encoding="utf-8"))

        raw_columns = data.get("columns") or {}
        if not raw_columns:
            raise TemplateError(f"Field mapping has no columns: {path}")

        return cls(
            sheet=data.get("sheet"),
            header_row=int(data.get("header_row", 1)),
            start_row=int(data.get("start_row", 2)),
            style_row=data.get("style_row"),
            columns=[ColumnSpec.parse(letter, raw) for letter, raw in raw_columns.items()],
        )

    @property
    def letters(self) -> list[str]:
        return [c.letter for c in self.columns]

    def required_report_fields(self) -> set[str]:
        """Names the VigiFlow client must supply for a row to be complete."""
        return {c.name for c in self.columns if c.source in ("field", "list")}


@dataclass
class RunContext:
    """Values that are the same for every row the run writes."""

    validation_date: date | None = None
    values: dict[str, Any] = field(default_factory=dict)

    def get(self, name: str) -> Any:
        if name == "validation_date" and self.validation_date is not None:
            return self.validation_date
        return self.values.get(name)


class ExcelTemplateWriter:
    """Copies the template once per batch and appends one row per report."""

    def __init__(
        self,
        template: Path,
        destination: Path,
        mapping: FieldMapping | None = None,
        start_number: int = 1,
        run: RunContext | None = None,
    ):
        self.template = Path(template)
        self.destination = Path(destination)
        self.mapping = mapping or FieldMapping.load()
        self.run = run or RunContext()
        self._next_number = start_number
        self._workbook = None
        self._sheet = None
        self._next_row = self.mapping.start_row
        self._styles: dict[str, Any] = {}
        self._written = 0

    # -- lifecycle ---------------------------------------------------------

    def __enter__(self) -> "ExcelTemplateWriter":
        self.open()
        return self

    def __exit__(self, exc_type, exc, tb) -> None:
        # Save even on failure: a partly filled workbook is evidence, and the
        # consumer only attaches it to the work item when the batch succeeded.
        self.close()

    def open(self) -> None:
        from openpyxl import load_workbook

        if not self.template.exists():
            raise TemplateError(
                f"Excel template not found: {self.template}. "
                "Regenerate it with scripts/make_template_from_sample.py."
            )
        self.destination.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(self.template, self.destination)

        self._workbook = load_workbook(self.destination)
        if self.mapping.sheet:
            if self.mapping.sheet not in self._workbook.sheetnames:
                raise TemplateError(
                    f"Sheet {self.mapping.sheet!r} is not in {self.template.name}. "
                    f"Available: {', '.join(self._workbook.sheetnames)}"
                )
            self._sheet = self._workbook[self.mapping.sheet]
        else:
            self._sheet = self._workbook.active

        self._capture_styles()
        self._next_row = self._first_free_row()

    def _capture_styles(self) -> None:
        """Remember the styling of the template's specimen data row.

        openpyxl gives rows that never existed no formatting at all, so
        without this every row past the template's own would lose its
        borders and date format.
        """
        if not self.mapping.style_row:
            return
        for spec in self.mapping.columns:
            cell = self._sheet[f"{spec.letter}{self.mapping.style_row}"]
            self._styles[spec.letter] = {
                "font": copy(cell.font),
                "border": copy(cell.border),
                "fill": copy(cell.fill),
                "alignment": copy(cell.alignment),
                "number_format": cell.number_format,
            }

    def _first_free_row(self) -> int:
        """Append below any rows the template already carries."""
        letters = self.mapping.letters
        row = self.mapping.start_row
        while row <= self._sheet.max_row:
            if all(self._sheet[f"{letter}{row}"].value in (None, "") for letter in letters):
                break
            row += 1
        return row

    # -- writing -----------------------------------------------------------

    def write_report(self, record: ReportRecord) -> int:
        """Write one report as a row. Returns the correlative number used."""
        if self._sheet is None:
            raise TemplateError("Writer is not open")

        number = self._next_number
        for spec in self.mapping.columns:
            cell = self._sheet[f"{spec.letter}{self._next_row}"]
            cell.value = _coerce(self._raw_value(spec, record, number), spec.type)
            self._apply_style(cell, spec.letter)

        LOGGER.info("Row %s: report %s as N. %s", self._next_row, record.report_id, number)
        self._next_row += 1
        self._next_number += 1
        self._written += 1
        return number

    def skip_number(self) -> int:
        """Consume a correlative without writing a row.

        Used when a report cannot be read. The number is burned so the rows
        that follow, here and in every later batch, keep the correlative the
        producer planned for them. The gap is deliberate and is recorded in
        the batch result.
        """
        number = self._next_number
        self._next_number += 1
        LOGGER.info("Correlative N. %s skipped, no row written", number)
        return number

    def _raw_value(self, spec: ColumnSpec, record: ReportRecord, number: int) -> Any:
        if spec.source == "sequence":
            return number
        if spec.source == "run":
            return self.run.get(spec.name)
        if spec.source == "list":
            values = record.fields.get(spec.name) or []
            if isinstance(values, str):
                # A single value, or a client that joined them with ';'.
                values = [part.strip() for part in values.split(";") if part.strip()]
            return values[spec.index] if spec.index < len(values) else None
        return record.fields.get(spec.name)

    def _apply_style(self, cell, letter: str) -> None:
        style = self._styles.get(letter)
        if not style:
            return
        cell.font = copy(style["font"])
        cell.border = copy(style["border"])
        cell.fill = copy(style["fill"])
        cell.alignment = copy(style["alignment"])
        cell.number_format = style["number_format"]

    def close(self) -> Path:
        if self._workbook is not None:
            self._workbook.save(self.destination)
            self._workbook.close()
            self._workbook = None
            LOGGER.info("Saved %s reports to %s", self._written, self.destination)
        return self.destination

    @property
    def written(self) -> int:
        return self._written

    @property
    def next_number(self) -> int:
        """The correlative the next row would take. Useful for chaining batches."""
        return self._next_number


# -- value conversion ------------------------------------------------------


def _coerce(value: Any, value_type: str) -> Any:
    """Turn a scraped string into what the spreadsheet expects.

    A value that cannot be converted is written through as text rather than
    dropped: a visibly odd cell is better than a silently empty one.
    """
    if value is None or (isinstance(value, str) and not value.strip()):
        return None

    if value_type == "date":
        return _as_date(value)
    if value_type == "digits":
        return _as_digits(value)
    if value_type == "int":
        return _as_int(value, fallback=value)
    if value_type == "auto":
        return _as_int(value, fallback=str(value).strip())
    return str(value).strip() if not isinstance(value, (int, float, date, datetime)) else value


def _as_date(value: Any) -> Any:
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    LOGGER.warning("Could not read %r as a date, writing it as text", text)
    return text


def _as_digits(value: Any) -> Any:
    """An identifier's digits, with any leading zeros kept.

    Column H holds the numeric part of the world-wide id. Most are nine digits
    with no leading zero and belong in the sheet as numbers, matching the file
    the business produces by hand. But ids like PE-PERULAB-00009 exist, and
    writing those as a number would turn 00009 into 9 and lose the identifier.
    So a value with a leading zero is written as text, which is the only way
    Excel will keep it, and everything else stays a number.
    """
    text = re.sub(r"\D", "", str(value))
    if not text:
        return str(value).strip() or None
    if len(text) > 1 and text.startswith("0"):
        return text
    return int(text)


def _as_int(value: Any, fallback: Any) -> Any:
    if isinstance(value, bool):
        return fallback
    if isinstance(value, int):
        return value
    text = str(value).strip().replace(" ", "")
    try:
        return int(text)
    except ValueError:
        return fallback
